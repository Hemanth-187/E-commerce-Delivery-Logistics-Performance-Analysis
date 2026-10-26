"""Build the final Week 2 KPI dataset from the unchanged source CSV."""

from pathlib import Path

import pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "data" / "raw" / "DataCoSupplyChainDataset.csv"
OUTPUT = ROOT / "reports" / "17_Final_KPI_Dataset.xlsx"

ORDER = "Order Id"
ITEM = "Order Item Id"
ORDER_STATUS = "Order Status"
DELIVERY_STATUS = "Delivery Status"
MODE = "Shipping Mode"
ACTUAL = "Days for shipping (real)"
SCHEDULED = "Days for shipment (scheduled)"
ORDER_DATE = "order date (DateOrders)"
SHIP_DATE = "shipping date (DateOrders)"
RISK = "Late_delivery_risk"

FIELDS = [
    ORDER, ITEM, ORDER_STATUS, DELIVERY_STATUS, MODE,
    ACTUAL, SCHEDULED, ORDER_DATE, SHIP_DATE, RISK,
]


def main():
    print("Reading source data...")
    df = pd.read_csv(SOURCE, encoding="latin1", usecols=FIELDS, low_memory=False)

    missing_columns = sorted(set(FIELDS) - set(df.columns))
    if missing_columns:
        raise ValueError(f"Required source fields missing: {missing_columns}")

    source_rows = len(df)
    missing_order_ids = int(df[ORDER].isna().sum())
    valid = df.dropna(subset=[ORDER]).copy()

    valid[ACTUAL] = pd.to_numeric(valid[ACTUAL], errors="coerce")
    valid[SCHEDULED] = pd.to_numeric(valid[SCHEDULED], errors="coerce")
    valid[RISK] = pd.to_numeric(valid[RISK], errors="coerce")
    valid[ORDER_DATE] = pd.to_datetime(valid[ORDER_DATE], errors="coerce")
    valid[SHIP_DATE] = pd.to_datetime(valid[SHIP_DATE], errors="coerce")

    group = valid.groupby(ORDER, sort=False, dropna=True)

    # Verify consistency before accepting one representative value per order.
    checked_fields = [
        ORDER_STATUS, DELIVERY_STATUS, MODE, ACTUAL, SCHEDULED,
        ORDER_DATE, SHIP_DATE, RISK,
    ]
    consistency = {
        field: int(group[field].nunique(dropna=True).gt(1).sum())
        for field in checked_fields
    }

    # No conflicting values are allowed in the final order-level dataset.
    if any(consistency.values()):
        raise ValueError(
            "Conflicting values found within orders; investigate before "
            f"building final dataset: {consistency}"
        )

    order_level = group.agg(
        Order_Status=(ORDER_STATUS, "first"),
        Delivery_Status=(DELIVERY_STATUS, "first"),
        Shipping_Mode=(MODE, "first"),
        Actual_Shipping_Days=(ACTUAL, "first"),
        Scheduled_Shipping_Days=(SCHEDULED, "first"),
        Late_Delivery_Risk=(RISK, "first"),
        Order_Date=(ORDER_DATE, "first"),
        Shipping_Date=(SHIP_DATE, "first"),
        Source_Item_Rows=(ITEM, "size"),
        Distinct_Order_Item_IDs=(ITEM, "nunique"),
    ).reset_index()

    order_level["Duration_Variance_Days"] = (
        order_level["Actual_Shipping_Days"]
        - order_level["Scheduled_Shipping_Days"]
    )
    order_level["Order_to_Shipment_Elapsed_Days"] = (
        order_level["Shipping_Date"] - order_level["Order_Date"]
    ).dt.total_seconds() / 86400

    delivery = order_level["Delivery_Status"].astype("string").str.strip()
    status = order_level["Order_Status"].astype("string").str.strip()

    total_orders = int(order_level[ORDER].nunique())
    canceled_delivery = int(delivery.eq("Shipping canceled").sum())
    eligible_non_canceled = total_orders - canceled_delivery
    canceled_order_status = int(status.eq("CANCELED").sum())
    on_time = int(delivery.eq("Shipping on time").sum())
    late = int(delivery.eq("Late delivery").sum())
    advance = int(delivery.eq("Advance shipping").sum())

    delivery_counts = (
        order_level["Delivery_Status"].fillna("Missing")
        .value_counts(dropna=False).rename_axis("Delivery Status")
        .reset_index(name="Distinct Orders")
    )
    delivery_counts["Share of All Orders (%)"] = (
        delivery_counts["Distinct Orders"] / total_orders * 100
    )
    delivery_counts["Share of Non-Canceled Orders (%)"] = (
        delivery_counts["Distinct Orders"] / eligible_non_canceled * 100
    )
    delivery_counts.loc[
        delivery_counts["Delivery Status"].eq("Shipping canceled"),
        "Share of Non-Canceled Orders (%)"
    ] = pd.NA

    order_status_counts = (
        order_level["Order_Status"].fillna("Missing")
        .value_counts(dropna=False).rename_axis("Order Status")
        .reset_index(name="Distinct Orders")
    )
    order_status_counts["Share of All Orders (%)"] = (
        order_status_counts["Distinct Orders"] / total_orders * 100
    )

    # KPI summary: definitions and limitations travel with the values.
    kpis = [
        ("Source rows", source_rows, "rows",
         "Rows in the raw CSV; rows represent order items.", "Validated count"),
        ("Distinct orders", total_orders, "orders",
         "Unique non-missing Order Id values.", "Validated count"),
        ("Distinct order item IDs", int(df[ITEM].nunique()), "items",
         "Unique non-missing Order Item Id values.", "Validated count"),
        ("Eligible non-canceled orders", eligible_non_canceled, "orders",
         "Distinct orders excluding Delivery Status = Shipping canceled.",
         "Eligibility rule must be confirmed by team"),
        ("Shipping-canceled orders", canceled_delivery, "orders",
         "Distinct orders with Delivery Status = Shipping canceled.",
         "Validated category count"),
        ("Orders with Order Status = CANCELED", canceled_order_status, "orders",
         "Distinct orders with Order Status = CANCELED.",
         "Separate status-based count; not interchangeable with shipping cancellation"),
        ("Shipping on time orders", on_time, "orders",
         "Distinct orders labelled Shipping on time.", "Validated category count"),
        ("Late delivery orders", late, "orders",
         "Distinct orders labelled Late delivery.", "Validated category count"),
        ("Advance shipping orders", advance, "orders",
         "Distinct orders labelled Advance shipping; kept separate.",
         "Validated category count; classification policy open"),
        ("Strict on-time status rate", on_time / eligible_non_canceled * 100,
         "%",
         "Shipping on time orders / non-canceled orders x 100.",
         "Provisional; denominator and status mapping require team approval"),
        ("Late delivery status share", late / eligible_non_canceled * 100,
         "%",
         "Late delivery orders / non-canceled orders x 100.",
         "Provisional; denominator requires team approval"),
        ("Advance shipping status share", advance / eligible_non_canceled * 100,
         "%",
         "Advance shipping orders / non-canceled orders x 100.",
         "Reported separately; do not silently count as on-time"),
        ("Order Status CANCELED share", canceled_order_status / total_orders * 100,
         "%",
         "Order Status = CANCELED orders / all distinct orders x 100.",
         "Status share only; official cancellation definition open"),
        ("Delivery Status Shipping canceled share",
         canceled_delivery / total_orders * 100, "%",
         "Shipping canceled orders / all distinct orders x 100.",
         "Status share only; official cancellation definition open"),
        ("Average actual shipping days", order_level["Actual_Shipping_Days"].mean(),
         "days", "Mean of Days for shipping (real) across order-level records.",
         "Field meaning follows source dictionary; not independently verified delivery timestamp"),
        ("Average scheduled shipping days",
         order_level["Scheduled_Shipping_Days"].mean(), "days",
         "Mean of Days for shipment (scheduled) across order-level records.",
         "Field meaning follows source dictionary"),
        ("Average duration variance",
         order_level["Duration_Variance_Days"].mean(), "days",
         "Mean of actual shipping days minus scheduled shipping days.",
         "Duration variance only; not timestamp-based delivery delay"),
        ("Average order-to-shipment elapsed time",
         order_level["Order_to_Shipment_Elapsed_Days"].mean(), "days",
         "Mean shipping timestamp minus order timestamp.",
         "Derived elapsed time; do not call warehouse processing time without further evidence"),
        ("Customer rating by delivery outcome", None, "unavailable",
         "No verified customer-rating field or order-level rating linkage was identified.",
         "Unavailable; do not infer"),
        ("Timestamp-based delivery delay", None, "unavailable",
         "No separate actual-delivery and promised-delivery timestamps were identified.",
         "Unavailable"),
    ]
    kpi_summary = pd.DataFrame(
        kpis, columns=["KPI", "Value", "Unit", "Formula / Definition", "Status / Limitation"]
    )

    # Shipping-mode comparisons, using distinct order-level records.
    mode_rows = []
    for mode_name, part in order_level.groupby("Shipping_Mode", dropna=False):
        d = part["Delivery_Status"].astype("string").str.strip()
        mode_total = len(part)
        mode_canceled = int(d.eq("Shipping canceled").sum())
        mode_eligible = mode_total - mode_canceled
        mode_rows.append({
            "Shipping Mode": mode_name,
            "Distinct Orders": mode_total,
            "Shipping Canceled Orders": mode_canceled,
            "Eligible Non-Canceled Orders": mode_eligible,
            "Shipping on time Orders": int(d.eq("Shipping on time").sum()),
            "Late Delivery Orders": int(d.eq("Late delivery").sum()),
            "Advance Shipping Orders": int(d.eq("Advance shipping").sum()),
            "Strict On-Time Status Rate (%)":
                (int(d.eq("Shipping on time").sum()) / mode_eligible * 100)
                if mode_eligible else None,
            "Average Actual Shipping Days": part["Actual_Shipping_Days"].mean(),
            "Average Scheduled Shipping Days": part["Scheduled_Shipping_Days"].mean(),
            "Average Duration Variance Days": part["Duration_Variance_Days"].mean(),
        })
    mode_summary = pd.DataFrame(mode_rows)

    checks = [
        ("Source row count", source_rows, 180519, source_rows == 180519,
         "Compared with the established source profile."),
        ("Distinct order count", total_orders, 65752, total_orders == 65752,
         "Compared with prior order-count validation."),
        ("Missing Order Id rows", missing_order_ids, 0, missing_order_ids == 0,
         "Missing IDs are excluded from order-level calculations."),
        ("Order-level rows reconcile", len(order_level), total_orders,
         len(order_level) == total_orders, "One row per distinct Order Id."),
        ("Distinct order item IDs", int(df[ITEM].nunique()), 180519,
         int(df[ITEM].nunique()) == 180519, "Compared with established source profile."),
        ("Delivery status counts reconcile",
         int(delivery_counts["Distinct Orders"].sum()), total_orders,
         int(delivery_counts["Distinct Orders"].sum()) == total_orders,
         "All delivery status categories including missing values."),
        ("Order status counts reconcile",
         int(order_status_counts["Distinct Orders"].sum()), total_orders,
         int(order_status_counts["Distinct Orders"].sum()) == total_orders,
         "All order status categories including missing values."),
        ("Eligible plus shipping-canceled reconcile",
         eligible_non_canceled + canceled_delivery, total_orders,
         eligible_non_canceled + canceled_delivery == total_orders,
         "Eligibility excludes Delivery Status = Shipping canceled only."),
        ("All checked order fields consistent",
         sum(consistency.values()), 0, sum(consistency.values()) == 0,
         "Count of per-order conflicting values across checked fields."),
        ("Order and shipping dates parse",
         int(order_level[["Order_Date", "Shipping_Date"]].isna().sum().sum()), 0,
         int(order_level[["Order_Date", "Shipping_Date"]].isna().sum().sum()) == 0,
         "Both timestamp fields must parse."),
        ("Duration fields numeric and present",
         int(order_level[["Actual_Shipping_Days", "Scheduled_Shipping_Days"]]
             .isna().sum().sum()), 0,
         int(order_level[["Actual_Shipping_Days", "Scheduled_Shipping_Days"]]
             .isna().sum().sum()) == 0, "Both duration fields must be numeric and present."),
        ("Late risk values restricted to 0 or 1",
         int((~order_level["Late_Delivery_Risk"].isin([0, 1])).sum()), 0,
         int((~order_level["Late_Delivery_Risk"].isin([0, 1])).sum()) == 0,
         "Checks supplied flag values; does not independently validate the flag."),
    ]
    validation = pd.DataFrame(
        checks, columns=["Check", "Actual", "Expected", "Passed", "Notes"]
    )

    definitions = pd.DataFrame([
        ("Source", "data/raw/DataCoSupplyChainDataset.csv",
         "Read-only input; script does not modify raw data."),
        ("Analysis unit", "Order Id",
         "Order-level counts use unique Order Id values; source rows are item rows."),
        ("Item count", "Order Item Id",
         "Reported separately from distinct order count."),
        ("Eligible non-canceled orders",
         "All distinct orders minus Delivery Status = Shipping canceled",
         "Provisional eligibility definition; confirm with the team."),
        ("Strict on-time status",
         "Delivery Status = Shipping on time",
         "Advance shipping remains a separate category."),
        ("Cancellation status",
         "Order Status = CANCELED versus Delivery Status = Shipping canceled",
         "Separate fields and counts; official cancellation definition remains open."),
        ("Duration variance",
         "Days for shipping (real) - Days for shipment (scheduled)",
         "Describe as duration variance, not timestamp-based delivery delay."),
        ("Order-to-shipment elapsed time",
         "Shipping date timestamp - order date timestamp",
         "Derived elapsed time; not automatically equivalent to warehouse processing time."),
        ("Late delivery risk", "Late_delivery_risk",
         "Provided source flag; retained as-is and not independently validated."),
        ("Customer ratings", "Unavailable",
         "No verified rating field/linkage identified in the source field inventory."),
        ("Privacy", "Exclude direct customer identifiers",
         "Final workbook excludes customer names, emails, passwords, and customer IDs."),
        ("Causation", "Descriptive analysis only",
         "Segment differences describe associations and do not prove causes."),
    ], columns=["Topic", "Rule", "Explanation"])

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(OUTPUT, engine="openpyxl") as writer:
        kpi_summary.to_excel(writer, sheet_name="KPI Summary", index=False)
        order_level.to_excel(writer, sheet_name="Order-Level KPI Dataset", index=False)
        delivery_counts.to_excel(writer, sheet_name="Delivery Status Breakdown", index=False)
        order_status_counts.to_excel(writer, sheet_name="Order Status Breakdown", index=False)
        mode_summary.to_excel(writer, sheet_name="Shipping Mode Comparison", index=False)
        validation.to_excel(writer, sheet_name="Validation Checks", index=False)
        definitions.to_excel(writer, sheet_name="Definitions and Limits", index=False)

        for ws in writer.book.worksheets:
            ws.freeze_panes = "A2"
            ws.auto_filter.ref = ws.dimensions
            for cell in ws[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="1F4E78")
                cell.alignment = Alignment(wrap_text=True, vertical="top")
            for cells in ws.columns:
                letter = cells[0].column_letter
                width = max(
                    len(str(cell.value)) if cell.value is not None else 0
                    for cell in cells[:200]
                )
                ws.column_dimensions[letter].width = min(max(width + 2, 12), 42)

    print(f"\nWorkbook created: {OUTPUT}")
    print(f"Source rows: {source_rows:,}")
    print(f"Distinct orders: {total_orders:,}")
    print(f"Eligible non-canceled orders: {eligible_non_canceled:,}")
    print(f"Shipping-canceled orders: {canceled_delivery:,}")
    print(f"Order Status = CANCELED: {canceled_order_status:,}")
    print("\nValidation results:")
    print(validation[["Check", "Actual", "Expected", "Passed"]].to_string(index=False))
    print(f"\nAll checks passed: {bool(validation['Passed'].all())}")


if __name__ == "__main__":
    main()
