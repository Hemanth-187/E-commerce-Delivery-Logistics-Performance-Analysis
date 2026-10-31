from pathlib import Path
import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.chart import LineChart, BarChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "data/raw/DataCoSupplyChainDataset.csv"
OUTPUT = ROOT / "reports/13_Delivery_Trend_Analysis.xlsx"

EXPECTED = {
    "total_orders": 65752,
    "eligible_delivered": 62897,
    "late": 36048,
    "on_time": 11722,
    "advance": 15127,
    "canceled": 2855,
}

# Read only the fields needed. Raw source remains unchanged.
df = pd.read_csv(
    SOURCE,
    encoding="latin1",
    usecols=["Order Id", "order date (DateOrders)", "Delivery Status"]
)

df["order date (DateOrders)"] = pd.to_datetime(
    df["order date (DateOrders)"], errors="coerce"
)

# Build one record per order and validate order-level consistency.
orders = df.groupby("Order Id", dropna=False).agg(
    order_date=("order date (DateOrders)", "first"),
    date_values=("order date (DateOrders)", "nunique"),
    delivery_status=("Delivery Status", "first"),
    status_values=("Delivery Status", "nunique"),
    source_rows=("Order Id", "size"),
).reset_index()

if orders["Order Id"].isna().any():
    raise ValueError("Missing Order Id values found.")

date_conflicts = int((orders["date_values"] > 1).sum())
status_conflicts = int((orders["status_values"] > 1).sum())

if date_conflicts or status_conflicts:
    raise ValueError(
        f"Order-level conflicts found: dates={date_conflicts}, "
        f"statuses={status_conflicts}. Investigate before analysis."
    )

if orders["order_date"].isna().any():
    raise ValueError("Missing or invalid order dates found.")

expected_statuses = {
    "Late delivery", "Shipping on time",
    "Advance shipping", "Shipping canceled"
}
observed_statuses = set(orders["delivery_status"].dropna().unique())

if observed_statuses != expected_statuses:
    raise ValueError(f"Unexpected delivery statuses: {observed_statuses}")

orders["Month"] = orders["order_date"].dt.to_period("M").astype(str)
orders["Year"] = orders["order_date"].dt.year

# Monthly distinct-order counts.
monthly = pd.crosstab(
    orders["Month"], orders["delivery_status"]
).reindex(columns=sorted(expected_statuses), fill_value=0)

monthly = monthly.rename(columns={
    "Late delivery": "Late Delivery Orders",
    "Shipping on time": "Strict On-Time Orders",
    "Advance shipping": "Advance Shipping Orders",
    "Shipping canceled": "Shipping Canceled Orders",
})

monthly["Total Distinct Orders"] = monthly.sum(axis=1)
monthly["Eligible Delivered Orders"] = (
    monthly["Total Distinct Orders"] -
    monthly["Shipping Canceled Orders"]
)

monthly["Late Delivery Rate"] = (
    monthly["Late Delivery Orders"] /
    monthly["Eligible Delivered Orders"]
)
monthly["Strict On-Time Rate"] = (
    monthly["Strict On-Time Orders"] /
    monthly["Eligible Delivered Orders"]
)
monthly["On-Time-or-Early Rate"] = (
    (monthly["Strict On-Time Orders"] +
     monthly["Advance Shipping Orders"]) /
    monthly["Eligible Delivered Orders"]
)

monthly.index.name = "Month"
monthly = monthly.reset_index()

# Source-row counts are included separately to distinguish rows from orders.
source_rows = (
    df.assign(Month=df["order date (DateOrders)"].dt.to_period("M").astype(str))
      .groupby("Month").size()
      .rename("Source Rows")
)
monthly["Source Rows"] = monthly["Month"].map(source_rows).fillna(0).astype(int)

# Yearly aggregation uses counts first, then calculates rates.
count_columns = [
    "Total Distinct Orders", "Eligible Delivered Orders",
    "Late Delivery Orders", "Strict On-Time Orders",
    "Advance Shipping Orders", "Shipping Canceled Orders",
    "Source Rows"
]
yearly = monthly.copy()
yearly["Year"] = yearly["Month"].str[:4]
yearly = yearly.groupby("Year")[count_columns].sum().reset_index()

yearly["Late Delivery Rate"] = (
    yearly["Late Delivery Orders"] / yearly["Eligible Delivered Orders"]
)
yearly["Strict On-Time Rate"] = (
    yearly["Strict On-Time Orders"] / yearly["Eligible Delivered Orders"]
)
yearly["On-Time-or-Early Rate"] = (
    (yearly["Strict On-Time Orders"] +
     yearly["Advance Shipping Orders"]) /
    yearly["Eligible Delivered Orders"]
)
yearly["Period Coverage"] = yearly["Year"].map({
    "2015": "Full calendar year",
    "2016": "Full calendar year",
    "2017": "Full calendar year in source date range; investigate volume change",
    "2018": "Partial year: January only",
}).fillna("Review coverage")

# Reconciliation checks.
totals = {
    "Total distinct orders": int(monthly["Total Distinct Orders"].sum()),
    "Eligible delivered orders": int(monthly["Eligible Delivered Orders"].sum()),
    "Late delivery orders": int(monthly["Late Delivery Orders"].sum()),
    "Strict on-time orders": int(monthly["Strict On-Time Orders"].sum()),
    "Advance shipping orders": int(monthly["Advance Shipping Orders"].sum()),
    "Shipping canceled orders": int(monthly["Shipping Canceled Orders"].sum()),
}

checks = [
    ("Distinct orders reconcile", totals["Total distinct orders"],
     EXPECTED["total_orders"]),
    ("Eligible delivered orders reconcile", totals["Eligible delivered orders"],
     EXPECTED["eligible_delivered"]),
    ("Late delivery orders reconcile", totals["Late delivery orders"],
     EXPECTED["late"]),
    ("Strict on-time orders reconcile", totals["Strict on-time orders"],
     EXPECTED["on_time"]),
    ("Advance shipping orders reconcile", totals["Advance shipping orders"],
     EXPECTED["advance"]),
    ("Shipping canceled orders reconcile", totals["Shipping canceled orders"],
     EXPECTED["canceled"]),
    ("Order date conflicts", date_conflicts, 0),
    ("Delivery status conflicts", status_conflicts, 0),
    ("Missing order IDs", int(df["Order Id"].isna().sum()), 0),
    ("Missing or invalid order dates", int(df["order date (DateOrders)"].isna().sum()), 0),
    ("Monthly periods", len(monthly), 37),
]
check_rows = []
for name, actual, expected in checks:
    check_rows.append([
        name, actual, expected,
        "PASS" if actual == expected else "FAIL"
    ])

if any(row[3] == "FAIL" for row in check_rows):
    raise ValueError("One or more reconciliation checks failed.")

# Create workbook.
wb = Workbook()
ws = wb.active
ws.title = "Monthly Delivery Trends"

def add_table(sheet, headers, rows):
    sheet.append(headers)
    for row in rows:
        sheet.append(list(row))
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions

monthly_headers = list(monthly.columns)
for _, row in monthly.iterrows():
    ws.append([row[col] for col in monthly_headers])
# Insert headers at top.
ws.insert_rows(1)
for col, header in enumerate(monthly_headers, start=1):
    ws.cell(1, col, header)

# Monthly rates sheet.
rate_ws = wb.create_sheet("Monthly Delivery Rates")
rate_headers = [
    "Month", "Eligible Delivered Orders", "Late Delivery Rate",
    "Strict On-Time Rate", "On-Time-or-Early Rate"
]
add_table(
    rate_ws, rate_headers,
    monthly[rate_headers].itertuples(index=False, name=None)
)

# Yearly summary.
year_ws = wb.create_sheet("Yearly Summary")
year_headers = list(yearly.columns)
add_table(
    year_ws, year_headers,
    yearly[year_headers].itertuples(index=False, name=None)
)

# Trend charts.
chart_ws = wb.create_sheet("Trend Charts")

volume_chart = BarChart()
volume_chart.title = "Monthly Distinct Orders by Delivery Status"
volume_chart.y_axis.title = "Orders"
volume_chart.x_axis.title = "Month"
volume_chart.style = 10
volume_chart.height = 9
volume_chart.width = 20

# Monthly Delivery Trends columns:
# A Month, B Advance, C Late, D Canceled, E On-time, then totals/rates.
for col in [2, 3, 4, 5]:
    data = Reference(ws, min_col=col, min_row=1, max_row=ws.max_row)
    volume_chart.add_data(data, titles_from_data=True)

categories = Reference(ws, min_col=1, min_row=2, max_row=ws.max_row)
volume_chart.set_categories(categories)
chart_ws.add_chart(volume_chart, "A1")

rate_chart = LineChart()
rate_chart.title = "Monthly Delivery Performance Rates"
rate_chart.y_axis.title = "Rate"
rate_chart.x_axis.title = "Month"
rate_chart.style = 13
rate_chart.height = 9
rate_chart.width = 20
rate_chart.y_axis.numFmt = "0%"
rate_chart.y_axis.scaling.min = 0
rate_chart.y_axis.scaling.max = 1

rate_data = Reference(rate_ws, min_col=3, max_col=5,
                      min_row=1, max_row=rate_ws.max_row)
rate_chart.add_data(rate_data, titles_from_data=True)
rate_chart.set_categories(
    Reference(rate_ws, min_col=1, min_row=2, max_row=rate_ws.max_row)
)
chart_ws.add_chart(rate_chart, "A20")

# Validation checks.
check_ws = wb.create_sheet("Validation Checks")
add_table(check_ws, ["Check", "Actual", "Expected", "Result"], check_rows)

# Definitions and limitations.
definitions = [
    ["Topic", "Definition / Limitation"],
    ["Source", "data/raw/DataCoSupplyChainDataset.csv; read-only."],
    ["Date basis", "Monthly grouping uses order date (DateOrders)."],
    ["Counting unit", "Distinct Order Id; source rows are reported separately."],
    ["Eligible delivered orders", "Total distinct orders minus Shipping canceled orders."],
    ["Late delivery rate", "Late delivery orders / eligible delivered orders."],
    ["Strict on-time rate", "Shipping on time orders / eligible delivered orders."],
    ["On-time-or-early rate", "(Shipping on time + Advance shipping) / eligible delivered orders."],
    ["Advance shipping", "Kept separate from strict on-time; included only in on-time-or-early."],
    ["Canceled shipments", "Shipping canceled excluded from delivered-order rate denominators."],
    ["Partial year", "2018 contains January only and must not be compared as a full year."],
    ["Coverage caution", "Source-row volume changes notably in late 2017; investigate coverage before interpreting trend changes."],
    ["Interpretation", "Descriptive trends do not establish causes of delivery delays."],
    ["Carrier limitation", "No carrier identifier is used in this Member 1 analysis."],
    ["Predictive modeling", "No predictive model is created in this workbook."],
    ["Data quality", "No rows are deleted or modified. Conflicting order dates/statuses cause the script to stop."],
]
def_ws = wb.create_sheet("Definitions and Limits")
add_table(def_ws, definitions[0], definitions[1:])

# Formatting.
for sheet in wb.worksheets:
    sheet.freeze_panes = "A2"
    sheet.sheet_view.showGridLines = False
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    sheet.row_dimensions[1].height = 30
    for column_cells in sheet.columns:
        letter = get_column_letter(column_cells[0].column)
        max_length = max(
            (len(str(cell.value)) if cell.value is not None else 0)
            for cell in column_cells
        )
        sheet.column_dimensions[letter].width = min(max(max_length + 2, 13), 48)
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

# Format percentage columns.
for sheet_name in ["Monthly Delivery Trends", "Monthly Delivery Rates", "Yearly Summary"]:
    sheet = wb[sheet_name]
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            if isinstance(cell.value, float) and (
                "Rate" in str(sheet.cell(1, cell.column).value)
            ):
                cell.number_format = "0.00%"

wb.save(OUTPUT)

# Reopen and verify workbook structure.
check_wb = load_workbook(OUTPUT, read_only=True)
expected_sheets = [
    "Monthly Delivery Trends", "Monthly Delivery Rates",
    "Yearly Summary", "Trend Charts",
    "Validation Checks", "Definitions and Limits"
]
if check_wb.sheetnames != expected_sheets:
    raise ValueError(f"Unexpected workbook sheets: {check_wb.sheetnames}")

print("WORKBOOK CREATED:", OUTPUT.relative_to(ROOT))
print("SHEETS:", check_wb.sheetnames)
print("MONTHLY PERIODS:", len(monthly))
print("TOTAL DISTINCT ORDERS:", totals["Total distinct orders"])
print("ELIGIBLE DELIVERED ORDERS:", totals["Eligible delivered orders"])
print("LATE DELIVERY RATE:", f'{totals["Late delivery orders"] / totals["Eligible delivered orders"]:.2%}')
print("STRICT ON-TIME RATE:", f'{totals["Strict on-time orders"] / totals["Eligible delivered orders"]:.2%}')
print("RECONCILIATION CHECKS:", len(check_rows))
print("FAILED CHECKS:", sum(row[3] == "FAIL" for row in check_rows))
print("SUCCESS: Workbook saved and reopened for structural verification.")
